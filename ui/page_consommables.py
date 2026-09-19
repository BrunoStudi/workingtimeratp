import json
import tkinter as tk
import tkinter.messagebox as msg
import tkinter.ttk as ttk

import customtkinter as ctk
from openpyxl import load_workbook
from tkinter import messagebox

from config.paths import CONSOMMABLES_FILE, EXCEL_RESEAU
from utils.page_lang import PageLang


TREEVIEW_STYLE = "Consommables.Treeview"
TREEVIEW_HEADING_STYLE = "Consommables.Treeview.Heading"
NETWORK_SHEET_NAME = "referentiel-articles"
INVALID_DESCRIPTIONS = {"~~null~~", "~null~", "null", "none"}

COLUMNS = (
    "id",
    "nom",
    "quantite_stoe",
    "quantite_vg",
)

# Assure que le dossier de stockage local existe.
CONSOMMABLES_FILE.parent.mkdir(parents=True, exist_ok=True)


# ---------------------------------------------------------------------------
# Stockage JSONL
# ---------------------------------------------------------------------------

def load_consommables():
    """Charge les consommables enregistrés localement."""
    if not CONSOMMABLES_FILE.exists():
        return []

    consommables = []

    with CONSOMMABLES_FILE.open("r", encoding="utf-8") as file:
        for line in file:
            if line.strip():
                consommables.append(json.loads(line))

    return consommables


def save_consommable(consommable):
    """Ajoute un consommable au fichier JSONL local."""
    with CONSOMMABLES_FILE.open("a", encoding="utf-8") as file:
        file.write(
            json.dumps(consommable, ensure_ascii=False) + "\n"
        )


def delete_from_jsonl(cid):
    """Supprime un consommable du fichier JSONL à partir de son identifiant."""
    consommables = load_consommables()
    target_id = str(cid).strip()

    consommables = [
        consommable
        for consommable in consommables
        if str(consommable["id"]).strip() != target_id
    ]

    _write_consommables(consommables)


def _write_consommables(consommables):
    """Réécrit entièrement le fichier JSONL."""
    with CONSOMMABLES_FILE.open("w", encoding="utf-8") as file:
        for consommable in consommables:
            file.write(
                json.dumps(consommable, ensure_ascii=False) + "\n"
            )


# ---------------------------------------------------------------------------
# Page
# ---------------------------------------------------------------------------

class PageConsommables(ctk.CTkFrame):
    """Gestion locale des consommables et synchronisation avec le référentiel EK1."""

    def __init__(self, parent, app):
        super().__init__(parent)

        self.app = app
        self.lang_util = PageLang(app)

        self._create_header()
        self._create_form()
        self._create_search_bar()
        self._create_table()
        self._create_context_menu()

        self.load_data()

    # ------------------------------------------------------------------
    # Construction de l'interface
    # ------------------------------------------------------------------

    def _create_header(self):
        header = ctk.CTkFrame(
            self,
            border_width=1,
            border_color="blue",
            fg_color="#1E5CC4",
            height=60,
        )
        header.pack(fill="x", padx=5, pady=5)
        header.pack_propagate(False)

        self.header_label = ctk.CTkLabel(
            header,
            text=self.lang_util.t("consommables_titre"),
            font=("Roboto", 24),
            text_color="white",
        )
        self.header_label.pack(expand=True)

    def _create_form(self):
        form_frame = ctk.CTkFrame(self)
        form_frame.pack(fill="x", padx=10, pady=5)

        self.entry_id = ctk.CTkEntry(
            form_frame,
            width=150,
            placeholder_text=self.lang_util.t("consommable_id"),
        )
        self.entry_id.pack(side="left", padx=5, pady=5)

        self.entry_nom = ctk.CTkEntry(
            form_frame,
            width=300,
            placeholder_text=self.lang_util.t("consommable_nom"),
        )
        self.entry_nom.pack(side="left", padx=5, pady=5)

        self.btn_add = ctk.CTkButton(
            form_frame,
            text=self.lang_util.t("ajouter"),
            command=self.add_consommable,
        )
        self.btn_add.pack(side="left", padx=5)

        self.btn_refresh_network = ctk.CTkButton(
            form_frame,
            text="Rafraîchir depuis le réseau",
            command=self.refresh_from_network,
        )
        self.btn_refresh_network.pack(pady=10)

    def _create_search_bar(self):
        search_frame = ctk.CTkFrame(self)
        search_frame.pack(fill="x", padx=10, pady=(5, 0))

        self.search_entry = ctk.CTkEntry(
            search_frame,
            placeholder_text=self.lang_util.t("rechercher"),
            width=480,
        )
        self.search_entry.pack(side="left", padx=5)

        self.search_btn = ctk.CTkButton(
            search_frame,
            text=self.lang_util.t("rechercher"),
            width=120,
            command=self.search_consommables,
        )
        self.search_btn.pack(side="left", padx=5, pady=5)

        # Réinitialise automatiquement la liste lorsque le champ est vidé
        # et applique le filtre pendant la saisie, comme dans la version actuelle.
        self.search_entry.bind(
            "<KeyRelease>",
            lambda event: self.load_data(self.search_entry.get()),
        )

    def _create_table(self):
        table_frame = ctk.CTkFrame(self)
        table_frame.pack(
            fill="both",
            expand=True,
            padx=10,
            pady=5,
        )

        self._configure_treeview_style()

        scrollbar_y = ttk.Scrollbar(
            table_frame,
            orient="vertical",
        )
        scrollbar_y.pack(
            side="right",
            fill="y",
        )

        self.tree = ttk.Treeview(
            table_frame,
            columns=COLUMNS,
            show="headings",
            style=TREEVIEW_STYLE,
            height=20,
            yscrollcommand=scrollbar_y.set,
        )

        self._configure_treeview_tags()
        self._configure_treeview_columns()

        self.tree.pack(
            side="left",
            fill="both",
            expand=True,
        )

        scrollbar_y.configure(
            command=self.tree.yview,
        )

    @staticmethod
    def _configure_treeview_style():
        style = ttk.Style()
        style.theme_use("default")

        style.configure(
            TREEVIEW_HEADING_STYLE,
            background="#DADADA",
            foreground="#000000",
            font=("Roboto", 10),
        )
        style.map(
            TREEVIEW_HEADING_STYLE,
            background=[("active", "#C8C8C8")],
        )

    def _configure_treeview_tags(self):
        self.tree.tag_configure(
            "rupture",
            background="#FF3030",
            foreground="black",
        )
        self.tree.tag_configure(
            "stoe_vide",
            background="#FFA500",
            foreground="black",
        )

    def _configure_treeview_columns(self):
        self.tree.heading(
            "id",
            text=self.lang_util.t("consommable_id"),
        )
        self.tree.heading(
            "nom",
            text=self.lang_util.t("consommable_nom"),
        )
        self.tree.heading(
            "quantite_stoe",
            text="Quantité STOE",
        )
        self.tree.heading(
            "quantite_vg",
            text="Quantité VG",
        )

        self.tree.column(
            "id",
            width=20,
            anchor="center",
        )
        self.tree.column(
            "nom",
            width=500,
            anchor="w",
        )
        self.tree.column(
            "quantite_stoe",
            width=20,
            anchor="center",
        )
        self.tree.column(
            "quantite_vg",
            width=20,
            anchor="center",
        )

    def _create_context_menu(self):
        self.menu = tk.Menu(
            self,
            tearoff=0,
        )
        self.menu.add_command(
            label=self.lang_util.t("supprimer"),
            command=self.delete_selected,
        )
        self.tree.bind(
            "<Button-3>",
            self.show_context_menu,
        )

    # ------------------------------------------------------------------
    # Menu contextuel / suppression
    # ------------------------------------------------------------------

    def show_context_menu(self, event):
        row_id = self.tree.identify_row(event.y)

        if not row_id:
            return

        self.tree.selection_set(row_id)
        self.menu.tk_popup(
            event.x_root,
            event.y_root,
        )

    def delete_selected(self):
        if not msg.askyesno(
            "Confirmation",
            "Supprimer ce consommable ?",
        ):
            return

        selected = self.tree.selection()

        if not selected:
            return

        item = selected[0]
        values = self.tree.item(item)["values"]
        cid = str(values[0]).strip()

        self.tree.delete(item)
        delete_from_jsonl(cid)

    # ------------------------------------------------------------------
    # Chargement / recherche / ajout
    # ------------------------------------------------------------------

    def load_data(self, search_text=""):
        self.tree.delete(
            *self.tree.get_children()
        )

        search_text = search_text.lower().strip()

        for item in load_consommables():
            cid = str(item["id"])
            nom = str(item["nom"])
            quantite_stoe = item.get("quantite_stoe", 0)
            quantite_vg = item.get("quantite_vg", 0)

            if (
                search_text
                and search_text not in cid.lower()
                and search_text not in nom.lower()
            ):
                continue

            self._insert_tree_item(
                cid=cid,
                nom=nom.upper(),
                quantite_stoe=quantite_stoe,
                quantite_vg=quantite_vg,
            )

    def consommable_exists(self, cid):
        cid = str(cid).strip()

        for row in self.tree.get_children():
            existing_id = str(
                self.tree.item(row)["values"][0]
            ).strip()

            if existing_id == cid:
                return True

        return False

    def add_consommable(self):
        cid = self.entry_id.get().strip()
        nom = self.entry_nom.get().strip().upper()

        if not cid or not nom:
            return

        if self.consommable_exists(cid):
            return

        consommable = {
            "id": cid,
            "nom": nom,
        }

        save_consommable(consommable)

        # Conserve exactement le comportement actuel : un consommable
        # ajouté manuellement n'a pas encore de quantités réseau.
        self.tree.insert(
            "",
            "end",
            values=(cid, nom),
            tags=("id_bold",),
        )

        self.entry_id.delete(0, "end")
        self.entry_nom.delete(0, "end")

    def search_consommables(self):
        self.load_data(
            self.search_entry.get()
        )

    # ------------------------------------------------------------------
    # Synchronisation réseau
    # ------------------------------------------------------------------

    def refresh_from_network(self):
        workbook = None

        try:
            workbook = load_workbook(
                EXCEL_RESEAU,
                read_only=True,
                data_only=True,
            )

            sheet = workbook[NETWORK_SHEET_NAME]
            articles_reseau = self._read_network_articles(sheet)
            consommables_locaux = self._load_local_consumables_by_id()

            nouveaux, modifies = self._merge_network_articles(
                consommables_locaux,
                articles_reseau,
            )

            _write_consommables(
                consommables_locaux.values()
            )

            self._refresh_tree_from_consumables(
                consommables_locaux.values()
            )

            messagebox.showinfo(
                "Rafraîchissement terminé",
                f"{len(articles_reseau)} articles trouvés dans le référentiel.\n\n"
                f"{nouveaux} nouveaux articles ajoutés.\n"
                f"{modifies} articles mis à jour.\n\n"
                f"{len(consommables_locaux)} consommables enregistrés localement.",
            )

        except FileNotFoundError:
            messagebox.showerror(
                "Erreur",
                "Le fichier eBOARD_AIC EK1.xlsx est introuvable.",
            )

        except KeyError:
            messagebox.showerror(
                "Erreur",
                "L'onglet 'referentiel-articles' est introuvable.",
            )

        except Exception as error:
            # Conservé comme dernier filet de sécurité pour les erreurs
            # provenant du fichier Excel, du partage réseau ou d'openpyxl.
            messagebox.showerror(
                "Erreur",
                f"Impossible de rafraîchir le référentiel :\n\n{error}",
            )

        finally:
            if workbook is not None:
                workbook.close()

    @staticmethod
    def _read_network_articles(sheet):
        articles_reseau = []

        for row in sheet.iter_rows(
            min_row=2,
            values_only=True,
        ):
            code_article = row[0]
            description = row[2]

            if code_article is None or description is None:
                continue

            description = str(description).strip()

            if not description:
                continue

            if description.lower() in INVALID_DESCRIPTIONS:
                continue

            code_article = str(code_article).strip()

            quantite_stoe = (
                row[6]
                if row[6] is not None
                else 0
            )
            quantite_vg = (
                row[7]
                if row[7] is not None
                else 0
            )

            articles_reseau.append(
                {
                    "id": code_article,
                    "nom": description,
                    "quantite_stoe": quantite_stoe,
                    "quantite_vg": quantite_vg,
                }
            )

        return articles_reseau

    @staticmethod
    def _load_local_consumables_by_id():
        consommables_locaux = {}

        if not CONSOMMABLES_FILE.exists():
            return consommables_locaux

        with CONSOMMABLES_FILE.open(
            "r",
            encoding="utf-8",
        ) as file:
            for line in file:
                line = line.strip()

                if not line:
                    continue

                try:
                    consommable = json.loads(line)
                    code = str(
                        consommable["id"]
                    ).strip()
                    consommables_locaux[code] = consommable

                except (json.JSONDecodeError, KeyError):
                    continue

        return consommables_locaux

    @staticmethod
    def _merge_network_articles(
        consommables_locaux,
        articles_reseau,
    ):
        nouveaux = 0
        modifies = 0

        for article in articles_reseau:
            code = article["id"]

            if code not in consommables_locaux:
                consommables_locaux[code] = article
                nouveaux += 1
                continue

            ancien = consommables_locaux[code]
            modification = False

            if ancien.get("nom", "") != article["nom"]:
                ancien["nom"] = article["nom"]
                modification = True

            if (
                ancien.get("quantite_stoe", 0)
                != article["quantite_stoe"]
            ):
                ancien["quantite_stoe"] = article["quantite_stoe"]
                modification = True

            if (
                ancien.get("quantite_vg", 0)
                != article["quantite_vg"]
            ):
                ancien["quantite_vg"] = article["quantite_vg"]
                modification = True

            if modification:
                modifies += 1

        return nouveaux, modifies

    def _refresh_tree_from_consumables(self, consommables):
        self.tree.delete(
            *self.tree.get_children()
        )

        for consommable in consommables:
            self._insert_tree_item(
                cid=consommable.get("id", ""),
                nom=consommable.get("nom", ""),
                quantite_stoe=consommable.get("quantite_stoe", 0),
                quantite_vg=consommable.get("quantite_vg", 0),
                uppercase_name=False,
            )

    def _insert_tree_item(
        self,
        cid,
        nom,
        quantite_stoe=0,
        quantite_vg=0,
        uppercase_name=True,
    ):
        tags = self._get_stock_tags(
            quantite_stoe,
            quantite_vg,
        )

        displayed_name = (
            str(nom).upper()
            if uppercase_name
            else str(nom)
        )

        self.tree.insert(
            "",
            "end",
            values=(
                cid,
                displayed_name,
                quantite_stoe,
                quantite_vg,
            ),
            tags=tags,
        )

    @staticmethod
    def _get_stock_tags(
        quantite_stoe,
        quantite_vg,
    ):
        if quantite_stoe == 0 and quantite_vg == 0:
            return ("id_bold", "rupture")

        if quantite_stoe == 0 and quantite_vg >= 1:
            return ("id_bold", "stoe_vide")

        return ("id_bold",)

    # ------------------------------------------------------------------
    # Rafraîchissement de la langue
    # ------------------------------------------------------------------

    def refresh_language(self):
        self.lang_util = PageLang(self.app)

        self.header_label.configure(
            text=self.lang_util.t("consommables_titre")
        )
        self.entry_id.configure(
            placeholder_text=self.lang_util.t("consommable_id")
        )
        self.entry_nom.configure(
            placeholder_text=self.lang_util.t("consommable_nom")
        )
        self.btn_add.configure(
            text=self.lang_util.t("ajouter")
        )
        self.tree.heading(
            "id",
            text=self.lang_util.t("consommable_id"),
        )
        self.tree.heading(
            "nom",
            text=self.lang_util.t("consommable_nom"),
        )
        self.search_entry.configure(
            placeholder_text=self.lang_util.t("rechercher")
        )
        self.search_btn.configure(
            text=self.lang_util.t("rechercher")
        )
